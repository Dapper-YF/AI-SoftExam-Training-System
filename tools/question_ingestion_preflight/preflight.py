from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import shutil
import sys
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable


SCHEMA_VERSION = "question-import/1.0"
RIGHTS_STATUS = "UNVERIFIED"
IMPORT_STATUS = "PREFLIGHT"
MAX_FILE_BYTES = 50 * 1024 * 1024

QUESTION_START = re.compile(
    r"^\s*(?:(?:试题\s*)[（(](?P<paren>\d{1,3})[）)]|"
    r"第\s*(?P<ordinal>\d{1,3})\s*题|"
    r"(?:试题\s*)?(?P<plain>\d{1,3})\s*[.．、:：])\s*"
)
OPTION_LINE = re.compile(r"^\s*([A-H])\s*[.．、)]\s*(.+)$", re.IGNORECASE)
ANSWER_LINE = re.compile(r"(?:参考答案|正确答案|答案|答)\s*[：:是为]?\s*([A-H]+|正确|错误|对|错|true|false|\d+(?:\.\d+)?)", re.IGNORECASE)
REFERENCE_ANSWER_LINE = re.compile(r"^\s*(?:参考答案|正确答案|答案|答)\s*[：:是为]?\s*(.+?)\s*$", re.IGNORECASE)
REFERENCE_SECTION_LINE = re.compile(r"^\s*[【\[]?\s*参考答案\s*[】\]]?\s*$")
EXPLANATION_LINE = re.compile(r"^\s*(?:答案解析|解析)\s*[：:]\s*(.*?)\s*$")
CASE_TITLE_LINE = re.compile(r"^\s*(?:案例|试题)\s*([一二三四五六七八九十百\d]+)\s*(?:[（(].*?[）)])?\s*$")
CASE_SUBQUESTION_LINE = re.compile(
    r"^\s*[\[【]?\s*(?:问题|小题)\s*(\d{1,3})\s*[\]】]?"
    r"\s*(?:[（(]\s*(\d+(?:\.\d+)?)\s*分\s*[）)])?\s*(.*)$"
)
ESSAY_TITLE_LINE = re.compile(r"^\s*(?:论文题(?:目)?|题目)\s*[：:]\s*(.+?)\s*$")
ESSAY_REQUIREMENT_LINE = re.compile(r"^\s*(?:写作要求|论文要求|要求)\s*[：:]\s*(.+?)\s*$")
ESSAY_REFERENCE_LINE = re.compile(r"^\s*(?:参考结构|参考答案|解析)\s*[：:]\s*(.+?)\s*$")
WORD_COUNT_RANGE = re.compile(r"(\d{3,5})\s*(?:至|到|[-~～—])\s*(\d{3,5})\s*字")
XLSX_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "sourceName": ("来源",),
    "qualificationName": ("科目", "资格"),
    "subjectName": ("科目名称", "题目类型", "考试科目"),
    "yearNo": ("年份", "年度"),
    "sourceQuestionId": ("题目ID", "题目编号"),
    "questionType": ("类型", "题型"),
    "stem": ("题目内容", "题目正文", "题干"),
    "optionsText": ("选项",),
    "answerText": ("答案", "正确答案"),
    "explanation": ("解析", "答案解析"),
    "chapter": ("章节",),
    "knowledgePoint": ("知识点",),
    "sourceUrl": ("URL", "链接"),
    "capturedAt": ("抓取时间", "采集时间"),
}


@dataclass(frozen=True)
class DocumentScope:
    qualification_code: str
    subject_code: str
    subject_type: str
    paper_type: str
    paper_kind: str
    year_no: int | None
    confidence: str
    notes: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class SegmentScope:
    segment_id: str
    qualification_code: str
    subject_code: str
    subject_type: str
    paper_type: str
    paper_kind: str
    year_no: int | None
    title_anchor: str
    start_page: int
    end_page: int
    confidence: str
    rule: str


@dataclass
class ExtractedPage:
    page_no: int
    text: str


@dataclass
class SourceDocument:
    path: Path
    sha256: str
    size_bytes: int
    pages: list[ExtractedPage]
    assets: list[dict[str, Any]]
    warnings: list[dict[str, str]]
    structured_questions: list[dict[str, Any]] = field(default_factory=list)
    structure_report: dict[str, Any] = field(default_factory=dict)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Offline question ingestion preflight")
    parser.add_argument("--input", action="append", required=True, help="Input file or directory")
    parser.add_argument("--output", required=True, help="Output evidence directory")
    parser.add_argument("--years", help="Inclusive year range, for example 2017-2021")
    parser.add_argument("--include-doc", action="store_true", help="Include legacy .doc conversion detection")
    parser.add_argument(
        "--xlsx-structure-only",
        action="store_true",
        help="Write an XLSX structure report instead of canonical question preflight outputs",
    )
    args = parser.parse_args(argv)

    inputs = [Path(value) for value in args.input]
    output = Path(args.output)
    year_filter = parse_year_filter(args.years)
    files = discover_files(inputs, year_filter, include_doc=args.include_doc)
    if args.xlsx_structure_only:
        if len(files) != 1 or files[0].suffix.lower() != ".xlsx":
            raise SystemExit("--xlsx-structure-only requires exactly one XLSX input")
        report = write_xlsx_structure_report(files[0], output)
        print(json.dumps({
            "output": str(output),
            "sheets": report["sheetCount"],
            "questions": report["questionCount"],
            "duplicates": report["duplicateCount"],
            "unmapped": report["unmappedCount"],
        }, ensure_ascii=False, sort_keys=True))
        return 0
    result = run_preflight(files, output)
    print(json.dumps({
        "output": str(output),
        "documents": result["manifest"]["documentCount"],
        "questions": result["manifest"]["questionCount"],
        "errors": result["manifest"]["errorCount"],
        "duplicates": result["manifest"]["duplicateCount"],
    }, ensure_ascii=False, sort_keys=True))
    return 0 if result["manifest"]["documentCount"] > 0 else 2


def parse_year_filter(value: str | None) -> set[int] | None:
    if not value:
        return None
    match = re.fullmatch(r"(\d{4})(?:-(\d{4}))?", value.strip())
    if not match:
        raise SystemExit("--years must be YYYY or YYYY-YYYY")
    start = int(match.group(1))
    end = int(match.group(2) or start)
    if end < start:
        raise SystemExit("--years end must be >= start")
    return set(range(start, end + 1))


def discover_files(inputs: Iterable[Path], years: set[int] | None, include_doc: bool) -> list[Path]:
    suffixes = {".pdf", ".docx", ".xlsx"}
    if include_doc:
        suffixes.add(".doc")
    files: list[Path] = []
    for item in inputs:
        if item.is_dir():
            candidates = sorted(path for path in item.rglob("*") if path.is_file())
            apply_year_filter = True
        elif item.is_file():
            candidates = [item]
            apply_year_filter = False
        else:
            continue
        for path in candidates:
            if path.suffix.lower() not in suffixes:
                continue
            if apply_year_filter and years is not None:
                year = infer_year(path.name)
                if year not in years:
                    continue
            files.append(path)
    return stable_unique(files)


def stable_unique(paths: list[Path]) -> list[Path]:
    seen: set[str] = set()
    unique: list[Path] = []
    for path in sorted(paths, key=lambda p: str(p).casefold()):
        key = str(path.resolve()).casefold()
        if key not in seen:
            seen.add(key)
            unique.append(path)
    return unique


def infer_year(name: str) -> int | None:
    match = re.search(r"(20\d{2})", name)
    return int(match.group(1)) if match else None


def run_preflight(files: list[Path], output: Path) -> dict[str, Any]:
    if output.exists() and any(output.iterdir()):
        raise FileExistsError(f"refusing to overwrite non-empty preflight output: {output}")
    output.mkdir(parents=True, exist_ok=True)
    records: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    unresolved: list[dict[str, Any]] = []
    assets: list[dict[str, Any]] = []
    documents: list[dict[str, Any]] = []

    for path in files:
        try:
            document = extract_document(path)
            scope = infer_scope(path)
            if document.structured_questions:
                scope = DocumentScope(
                    "ROW_SCOPED",
                    "ROW_SCOPED",
                    "ROW_SCOPED",
                    "ROW_SCOPED",
                    "ROW_SCOPED",
                    None,
                    "HIGH",
                )
            split_plan = build_split_plan(document, scope)
            doc_errors = document.warnings + split_plan_errors(path, scope, split_plan)
            question_records = build_question_records(document, scope, split_plan)
            for record in question_records:
                validation_errors = validate_record(record)
                record["quality"]["validationErrors"] = validation_errors
                record["quality"]["importEligible"] = not validation_errors
                record["quality"]["needsManualReview"] = bool(
                    record["quality"].get("needsManualReview") or validation_errors
                )
                records.append(record)
                unresolved.extend(unresolved_rows_for_record(record, validation_errors))
            for warning in doc_errors:
                errors.append(error_row(path, document.sha256, warning["code"], warning["message"], warning.get("locator")))
                unresolved.append(unresolved_row(
                    path,
                    document.sha256,
                    warning.get("locator"),
                    "",
                    warning["code"],
                    "source",
                    warning["message"],
                ))
            for asset in document.assets:
                assets.append(asset)
            documents.append(document_manifest(path, document, scope, split_plan, len(question_records), len(doc_errors)))
        except Exception as exc:  # noqa: BLE001 - preflight must report per-file failures.
            sha = file_sha256(path) if path.exists() and path.is_file() else ""
            errors.append(error_row(path, sha, "DOCUMENT_PRECHECK_FAILED", str(exc), None))
            unresolved.append(unresolved_row(
                path, sha, None, "", "DOCUMENT_PRECHECK_FAILED", "source", str(exc)
            ))
            documents.append({
                "path": str(path),
                "sha256": sha,
                "sizeBytes": path.stat().st_size if path.exists() else 0,
                "status": "FAILED",
                "questionCount": 0,
                "errorCount": 1,
            })

    duplicates = duplicate_rows(records)
    record_by_external_key = {record["question"]["externalKey"]: record for record in records}
    for duplicate in duplicates:
        duplicate_record = record_by_external_key[duplicate["duplicateExternalKey"]]
        duplicate_record["quality"]["importEligible"] = False
        duplicate_record["quality"]["needsManualReview"] = True
        duplicate_record["quality"]["validationErrors"].append({
            "code": "DUPLICATE_CONTENT",
            "field": "question.contentHash",
            "message": f"duplicates {duplicate['firstExternalKey']}",
        })
        unresolved.append(unresolved_row(
            Path(duplicate_record["sourceTrace"]["sourcePath"]),
            duplicate_record["sourceTrace"]["documentSha256"],
            duplicate_record["sourceTrace"]["locator"],
            duplicate["duplicateExternalKey"],
            "DUPLICATE_CONTENT",
            "question.contentHash",
            f"duplicates {duplicate['firstExternalKey']}",
        ))
    unresolved = stable_unique_unresolved(unresolved)
    write_jsonl(output / "canonical-questions.jsonl", records)
    write_csv(output / "errors.csv", errors, ["sourcePath", "documentSha256", "code", "message", "locator"])
    write_csv(output / "duplicates.csv", duplicates, ["contentHash", "firstExternalKey", "duplicateExternalKey", "firstSource", "duplicateSource"])
    write_json(output / "assets.json", {"schemaVersion": "question-assets/1.0", "assets": assets})
    write_csv(
        output / "unresolved-items.csv",
        unresolved,
        ["sourcePath", "documentSha256", "locator", "externalKey", "code", "field", "message", "blocking"],
    )
    manifest = build_manifest(files, documents, records, errors, duplicates, assets, unresolved)
    write_json(output / "manifest.json", manifest)
    write_quality_report(output / "quality-report.md", manifest, documents, errors)
    return {
        "manifest": manifest,
        "records": records,
        "errors": errors,
        "duplicates": duplicates,
        "assets": assets,
        "unresolved": unresolved,
    }


def extract_document(path: Path) -> SourceDocument:
    size = path.stat().st_size
    sha = file_sha256(path)
    if size > MAX_FILE_BYTES:
        return SourceDocument(path, sha, size, [], [], [{"code": "FILE_TOO_LARGE", "message": f"file exceeds {MAX_FILE_BYTES} bytes"}])
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path, sha, size)
    if suffix == ".docx":
        return extract_docx(path, sha, size)
    if suffix == ".xlsx":
        return extract_xlsx(path, sha, size)
    if suffix == ".doc":
        return detect_legacy_doc(path, sha, size)
    return SourceDocument(path, sha, size, [], [], [{"code": "UNSUPPORTED_FILE_TYPE", "message": suffix}])


def extract_pdf(path: Path, sha: str, size: int) -> SourceDocument:
    try:
        from pypdf import PdfReader
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError("pypdf is required for PDF preflight") from exc

    reader = PdfReader(str(path))
    pages: list[ExtractedPage] = []
    assets: list[dict[str, Any]] = []
    warnings: list[dict[str, str]] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append(ExtractedPage(index, text))
        image_count = count_pdf_images(page)
        for image_index in range(1, image_count + 1):
            assets.append(asset_row(path, sha, f"page-{index}-image-{image_index}", "PDF_IMAGE", index, "page-xobject", ""))
        if not text.strip() and image_count:
            warnings.append({"code": "PDF_PAGE_TEXT_EMPTY_WITH_IMAGES", "message": "page may require OCR", "locator": f"page:{index}"})
    if not any(page.text.strip() for page in pages):
        warnings.append({"code": "PDF_TEXT_EXTRACTION_EMPTY", "message": "document may be scanned or encrypted", "locator": None})
    return SourceDocument(path, sha, size, pages, assets, warnings)


def count_pdf_images(page: Any) -> int:
    try:
        resources = page.get("/Resources") or {}
        xobjects = resources.get("/XObject") or {}
        count = 0
        for obj in xobjects.values():
            resolved = obj.get_object()
            if resolved.get("/Subtype") == "/Image":
                count += 1
        return count
    except Exception:
        return 0


def extract_docx(path: Path, sha: str, size: int) -> SourceDocument:
    try:
        from docx import Document
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError("python-docx is required for DOCX preflight") from exc

    document = Document(str(path))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            lines.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    assets: list[dict[str, Any]] = []
    with zipfile.ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            if name.startswith("word/media/"):
                info = archive.getinfo(name)
                assets.append(asset_row(path, sha, name, "DOCX_MEDIA", None, "word/media", str(info.file_size)))
    warnings = [] if lines else [{"code": "DOCX_TEXT_EXTRACTION_EMPTY", "message": "no visible text extracted", "locator": None}]
    return SourceDocument(path, sha, size, [ExtractedPage(1, "\n".join(lines))], assets, warnings)


def extract_xlsx(path: Path, sha: str, size: int) -> SourceDocument:
    try:
        import openpyxl
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError("openpyxl is required for XLSX preflight") from exc

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    pages: list[ExtractedPage] = []
    warnings: list[dict[str, str]] = []
    structured_questions: list[dict[str, Any]] = []
    for index, sheet_name in enumerate(workbook.sheetnames, start=1):
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        header_values = next(rows, None)
        headers = [stringify_cell(value) for value in header_values or []]
        header_map = {header: position for position, header in enumerate(headers) if header}
        if is_structured_question_sheet(header_map):
            for row_no, values in enumerate(rows, start=2):
                row = structured_xlsx_row(sheet_name, index, row_no, header_map, values)
                if row is not None:
                    structured_questions.append(row)
            pages.append(ExtractedPage(index, f"## sheet:{sheet_name}\n## structured-question-rows"))
        else:
            lines: list[str] = [f"## sheet:{sheet_name}", " | ".join(headers)]
            for values in rows:
                cells = [stringify_cell(value) for value in values]
                if any(cells):
                    lines.append(" | ".join(cells))
            pages.append(ExtractedPage(index, "\n".join(lines)))
    if not pages:
        warnings.append({"code": "XLSX_WORKBOOK_EMPTY", "message": "workbook has no sheets", "locator": None})
    report = build_xlsx_structure_report(path)
    return SourceDocument(path, sha, size, pages, [], warnings, structured_questions, report)


def is_structured_question_sheet(header_map: dict[str, int]) -> bool:
    return all(
        xlsx_column_position(header_map, key) is not None
        for key in ("qualificationName", "subjectName", "sourceQuestionId", "questionType", "stem", "answerText")
    )


def xlsx_column_position(header_map: dict[str, int], canonical_name: str) -> int | None:
    for alias in XLSX_HEADER_ALIASES[canonical_name]:
        if alias in header_map:
            return header_map[alias]
    return None


def structured_xlsx_row(
    sheet_name: str,
    sheet_index: int,
    row_no: int,
    header_map: dict[str, int],
    values: tuple[Any, ...],
) -> dict[str, Any] | None:
    def value(canonical_name: str) -> str:
        position = xlsx_column_position(header_map, canonical_name)
        return stringify_cell(values[position]) if position is not None and position < len(values) else ""

    stem = value("stem")
    if not stem:
        return None
    qualification_code = map_qualification(value("qualificationName"))
    subject_code, subject_type = map_subject(value("subjectName"), qualification_code)
    paper_type, paper_kind = map_paper_kind(f"{value('sourceName')} {value('subjectName')}")
    return {
        "sheetName": sheet_name,
        "sheetIndex": sheet_index,
        "rowNo": row_no,
        "sourceName": value("sourceName"),
        "qualificationName": value("qualificationName"),
        "qualificationCode": qualification_code,
        "subjectName": value("subjectName"),
        "subjectCode": subject_code,
        "subjectType": subject_type,
        "paperType": paper_type,
        "paperKind": paper_kind,
        "yearNo": normalize_year(value("yearNo"), value("sourceName")),
        "sourceQuestionId": value("sourceQuestionId"),
        "questionType": map_question_type(value("questionType")),
        "stem": stem,
        "optionsText": value("optionsText"),
        "answerText": value("answerText"),
        "explanation": value("explanation"),
        "chapter": value("chapter"),
        "knowledgePoint": value("knowledgePoint"),
        "sourceUrl": value("sourceUrl"),
        "capturedAt": value("capturedAt"),
    }


def map_qualification(value: str) -> str:
    normalized = value.replace(" ", "")
    if "信息系统项目管理师" in normalized or normalized == "高项":
        return "HIGH_PM"
    if "数据库系统工程师" in normalized:
        return "MIDDLE_DATABASE_SYSTEM_ENGINEER"
    return "UNKNOWN"


def map_subject(value: str, qualification_code: str) -> tuple[str, str]:
    normalized = value.replace(" ", "")
    if "论文" in normalized:
        return "ESSAY", "ESSAY"
    if "案例" in normalized:
        return ("CASE_ANALYSIS", "CASE") if qualification_code == "HIGH_PM" else ("APPLIED_TECHNOLOGY", "APPLICATION")
    if "应用技术" in normalized or "下午" in normalized:
        return "APPLIED_TECHNOLOGY", "APPLICATION"
    if "综合知识" in normalized or "基础知识" in normalized or "上午" in normalized:
        subject_code = (
            "BASIC_KNOWLEDGE"
            if qualification_code == "MIDDLE_DATABASE_SYSTEM_ENGINEER"
            else "COMPREHENSIVE_KNOWLEDGE"
        )
        return subject_code, "COMPREHENSIVE"
    return "UNKNOWN", "UNKNOWN"


def map_paper_kind(value: str) -> tuple[str, str]:
    normalized = value.replace(" ", "")
    if "模拟" in normalized or "一模" in normalized or "二模" in normalized:
        return "MOCK", "MOCK_PAPER"
    if "真题" in normalized or "历年" in normalized:
        return "REAL", "REAL_PAPER"
    return "UNKNOWN", "UNKNOWN"


def normalize_year(*values: str) -> int | None:
    for value in values:
        match = re.search(r"(?:19|20)\d{2}", value)
        if match:
            return int(match.group(0))
    return None


def map_question_type(value: str) -> str:
    normalized = value.strip().lower().replace("-", "_").replace(" ", "_")
    mappings = {
        "single_choice": "SINGLE_CHOICE",
        "single": "SINGLE_CHOICE",
        "单选题": "SINGLE_CHOICE",
        "multiple_choice": "MULTI_CHOICE",
        "multi_choice": "MULTI_CHOICE",
        "multiple": "MULTI_CHOICE",
        "多选题": "MULTI_CHOICE",
        "judge": "JUDGE",
        "judgement": "JUDGE",
        "true_false": "JUDGE",
        "判断题": "JUDGE",
        "case": "CASE",
        "案例题": "CASE",
        "essay": "ESSAY",
        "论文题": "ESSAY",
    }
    return mappings.get(normalized, "UNKNOWN")


def build_xlsx_structure_report(path: Path) -> dict[str, Any]:
    try:
        import openpyxl
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError("openpyxl is required for XLSX structure reporting") from exc

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets: list[dict[str, Any]] = []
    qualification_counts: Counter[str] = Counter()
    subject_counts: Counter[str] = Counter()
    question_type_counts: Counter[str] = Counter()
    year_counts: Counter[str] = Counter()
    paper_kind_counts: Counter[str] = Counter()
    content_hash_counts: Counter[str] = Counter()
    question_count = 0
    with_answer_count = 0
    with_explanation_count = 0
    unmapped_count = 0

    for sheet_index, sheet_name in enumerate(workbook.sheetnames, start=1):
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        header_values = next(rows, None)
        headers = [stringify_cell(value) for value in header_values or []]
        header_map = {header: position for position, header in enumerate(headers) if header}
        populated_rows = 0
        for row_no, values in enumerate(rows, start=2):
            if any(stringify_cell(value) for value in values):
                populated_rows += 1
            if not is_structured_question_sheet(header_map):
                continue
            row = structured_xlsx_row(sheet_name, sheet_index, row_no, header_map, values)
            if row is None:
                continue
            question_count += 1
            qualification_counts[row["qualificationCode"]] += 1
            subject_counts[row["subjectCode"]] += 1
            question_type_counts[row["questionType"]] += 1
            year_counts[str(row["yearNo"] or "UNKNOWN")] += 1
            paper_kind_counts[row["paperKind"]] += 1
            with_answer_count += int(bool(row["answerText"]))
            with_explanation_count += int(bool(row["explanation"]))
            unmapped_count += int(
                row["qualificationCode"] == "UNKNOWN"
                or row["subjectCode"] == "UNKNOWN"
                or row["questionType"] == "UNKNOWN"
            )
            fingerprint = sha256_text(json.dumps({
                "stem": clean_line(row["stem"]),
                "options": clean_line(row["optionsText"]),
                "answer": clean_line(row["answerText"]),
            }, ensure_ascii=False, sort_keys=True, separators=(",", ":")))
            content_hash_counts[fingerprint] += 1
        sheets.append({
            "name": sheet_name,
            "index": sheet_index,
            "state": sheet.sheet_state,
            "rowCount": populated_rows,
            "columnCount": len(headers),
            "headers": headers,
            "columnMapping": {
                header: {"columnIndex": position + 1, "header": header}
                for header, position in header_map.items()
            },
            "structuredQuestionSheet": is_structured_question_sheet(header_map),
        })
    return {
        "schemaVersion": "question-xlsx-structure-report/1.0",
        "sourceFile": str(path),
        "sourceSha256": file_sha256(path),
        "sheetCount": len(sheets),
        "sheets": sheets,
        "questionCount": question_count,
        "questionTypeDistribution": dict(sorted(question_type_counts.items())),
        "qualificationDistribution": dict(sorted(qualification_counts.items())),
        "subjectDistribution": dict(sorted(subject_counts.items())),
        "yearDistribution": dict(sorted(year_counts.items())),
        "paperKindDistribution": dict(sorted(paper_kind_counts.items())),
        "withAnswerCount": with_answer_count,
        "withExplanationCount": with_explanation_count,
        "duplicateCount": sum(count - 1 for count in content_hash_counts.values() if count > 1),
        "unmappedCount": unmapped_count,
        "nineSheetTemplateCompatible": len(sheets) == 9,
        "sixtyOneColumnTemplateCompatible": any(sheet["columnCount"] == 61 for sheet in sheets),
    }


def write_xlsx_structure_report(path: Path, output: Path) -> dict[str, Any]:
    if output.exists() and any(output.iterdir()):
        raise FileExistsError(f"refusing to overwrite non-empty XLSX report output: {output}")
    output.mkdir(parents=True, exist_ok=True)
    report = build_xlsx_structure_report(path)
    write_json(output / "xlsx-structure-report.json", report)
    lines = [
        "# XLSX Question Source Structure Report",
        "",
        f"- Source file: `{report['sourceFile']}`",
        f"- Source SHA-256: `{report['sourceSha256']}`",
        f"- Sheets: `{report['sheetCount']}`",
        f"- Questions: `{report['questionCount']}`",
        f"- With answer: `{report['withAnswerCount']}`",
        f"- With explanation: `{report['withExplanationCount']}`",
        f"- Duplicates: `{report['duplicateCount']}`",
        f"- Unmapped: `{report['unmappedCount']}`",
        f"- Nine-Sheet template compatible: `{report['nineSheetTemplateCompatible']}`",
        f"- 61-column template compatible: `{report['sixtyOneColumnTemplateCompatible']}`",
        "",
        "## Sheets and columns",
        "",
    ]
    for sheet in report["sheets"]:
        lines.append(
            f"- `{sheet['name']}`: rows=`{sheet['rowCount']}`, columns=`{sheet['columnCount']}`, "
            f"structured=`{sheet['structuredQuestionSheet']}`"
        )
        lines.append(f"  - Headers: `{', '.join(sheet['headers'])}`")
    for title, key in (
        ("Question type distribution", "questionTypeDistribution"),
        ("Qualification distribution", "qualificationDistribution"),
        ("Subject distribution", "subjectDistribution"),
        ("Year distribution", "yearDistribution"),
        ("Paper kind distribution", "paperKindDistribution"),
    ):
        lines.extend(["", f"## {title}", ""])
        for value, count in report[key].items():
            lines.append(f"- `{value}`: `{count}`")
    (output / "xlsx-structure-report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return report


def detect_legacy_doc(path: Path, sha: str, size: int) -> SourceDocument:
    converter = shutil.which("soffice") or shutil.which("libreoffice")
    if converter:
        warning = {"code": "DOC_CONVERTER_AVAILABLE", "message": f"legacy DOC conversion available via {converter}", "locator": None}
    else:
        warning = {"code": "DOC_CONVERTER_UNAVAILABLE", "message": "legacy DOC requires soffice/libreoffice before parsing", "locator": None}
    return SourceDocument(path, sha, size, [], [], [warning])


def stringify_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def infer_scope(path: Path) -> DocumentScope:
    name = path.name
    year = infer_year(name)
    normalized = name.replace(" ", "")
    if "数据库系统工程师" in normalized:
        qualification = "MIDDLE_DATABASE_SYSTEM_ENGINEER"
        if "上午" in normalized and "下午" not in normalized:
            return DocumentScope(qualification, "BASIC_KNOWLEDGE", "COMPREHENSIVE", "REAL", "REAL_PAPER", year, "HIGH")
        if "下午" in normalized and "上午" not in normalized:
            return DocumentScope(qualification, "APPLIED_TECHNOLOGY", "APPLICATION", "REAL", "REAL_PAPER", year, "HIGH")
        return DocumentScope(qualification, "MIXED_DATABASE_SUBJECTS", "MIXED", "REAL", "REAL_PAPER", year, "LOW", ["DATABASE_MIXED_SUBJECT_SOURCE"])
    if "信息系统项目管理师" in normalized or "高项" in normalized:
        paper = "MOCK" if ("一模" in normalized or "模拟" in normalized) else "REAL"
        kind = "MOCK_PAPER" if paper == "MOCK" else "REAL_PAPER"
        if "论文" in normalized:
            subject, subject_type = "ESSAY", "ESSAY"
        elif "案例" in normalized:
            subject, subject_type = "CASE_ANALYSIS", "CASE"
        elif "一模" in normalized or "试卷" in normalized:
            subject, subject_type = "MIXED_ADVANCED_PM_SUBJECTS", "MIXED"
        else:
            subject, subject_type = "COMPREHENSIVE_KNOWLEDGE", "COMPREHENSIVE"
        confidence = "MEDIUM" if subject_type == "MIXED" else "HIGH"
        notes = ["ADVANCED_PM_MIXED_SUBJECT_SOURCE"] if subject_type == "MIXED" else []
        return DocumentScope("HIGH_PM", subject, subject_type, paper, kind, year, confidence, notes)
    return DocumentScope("UNKNOWN", "UNKNOWN", "UNKNOWN", "UNKNOWN", "UNKNOWN", year, "LOW", ["SOURCE_SCOPE_UNKNOWN"])


def build_split_plan(document: SourceDocument, scope: DocumentScope) -> dict[str, Any]:
    if scope.subject_code == "MIXED_DATABASE_SUBJECTS":
        return mixed_database_split_plan(document, scope)
    if scope.subject_code == "MIXED_ADVANCED_PM_SUBJECTS":
        return mixed_advanced_pm_split_plan(document, scope)
    segment = SegmentScope(
        segment_id="S01",
        qualification_code=scope.qualification_code,
        subject_code=scope.subject_code,
        subject_type=scope.subject_type,
        paper_type=scope.paper_type,
        paper_kind=scope.paper_kind,
        year_no=scope.year_no,
        title_anchor="WHOLE_DOCUMENT",
        start_page=document.pages[0].page_no if document.pages else 1,
        end_page=document.pages[-1].page_no if document.pages else 1,
        confidence=scope.confidence,
        rule="single-scope-source",
    )
    return {"status": "WHOLE_DOCUMENT", "segments": [segment_to_dict(segment)], "errors": []}


def mixed_database_split_plan(document: SourceDocument, scope: DocumentScope) -> dict[str, Any]:
    rules = [
        ("BASIC_KNOWLEDGE", "COMPREHENSIVE", "DATABASE_AM", re.compile(r"(上午.*(综合|试卷|真题)|综合.*试卷)")),
        ("APPLIED_TECHNOLOGY", "APPLICATION", "DATABASE_PM", re.compile(r"(下午.*(案例|应用|试卷|真题)|应用技术)")),
    ]
    anchors = find_ordered_anchors(document.pages, rules)
    if not anchors or [anchor["rule"] for anchor in anchors] != ["DATABASE_AM", "DATABASE_PM"]:
        return failed_split_plan(
            "DATABASE_MIXED_SUBJECT_SOURCE",
            "database mixed source requires ordered AM and PM title anchors",
            scan_candidate_anchors(document.pages, rules),
            ["DATABASE_AM", "DATABASE_PM"],
        )
    return segmented_plan(document, scope, anchors)


def mixed_advanced_pm_split_plan(document: SourceDocument, scope: DocumentScope) -> dict[str, Any]:
    rules = [
        ("COMPREHENSIVE_KNOWLEDGE", "COMPREHENSIVE", "HIGH_PM_COMPREHENSIVE", re.compile(r"(综\s*合\s*知\s*识|上午.*(综合|试卷)|选\s*择\s*题)")),
        ("CASE_ANALYSIS", "CASE", "HIGH_PM_CASE", re.compile(r"(案\s*例\s*分\s*析|案\s*例\s*题|下午.*案例)")),
        ("ESSAY", "ESSAY", "HIGH_PM_ESSAY", re.compile(r"(论\s*文|论\s*述\s*题)")),
    ]
    anchors = find_ordered_anchors(document.pages, rules)
    if not anchors or [anchor["rule"] for anchor in anchors] != ["HIGH_PM_COMPREHENSIVE", "HIGH_PM_CASE", "HIGH_PM_ESSAY"]:
        return failed_split_plan(
            "ADVANCED_PM_MIXED_SUBJECT_SOURCE",
            "advanced PM mixed source requires ordered comprehensive/case/essay title anchors",
            scan_candidate_anchors(document.pages, rules),
            ["HIGH_PM_COMPREHENSIVE", "HIGH_PM_CASE", "HIGH_PM_ESSAY"],
        )
    return segmented_plan(document, scope, anchors)


def find_ordered_anchors(pages: list[ExtractedPage], rules: list[tuple[str, str, str, re.Pattern[str]]]) -> list[dict[str, Any]] | None:
    anchors: list[dict[str, Any]] = []
    next_rule = 0
    for page in pages:
        lines = [clean_line(line) for line in page.text.splitlines() if clean_line(line)]
        for line_no, line in enumerate(lines[:80]):
            if next_rule >= len(rules):
                break
            subject_code, subject_type, rule_name, pattern = rules[next_rule]
            if is_title_anchor_line(line) and pattern.search(line):
                anchors.append({
                    "pageNo": page.page_no,
                    "lineNo": line_no,
                    "titleAnchor": title_anchor(line),
                    "subjectCode": subject_code,
                    "subjectType": subject_type,
                    "rule": rule_name,
                })
                next_rule += 1
        if next_rule >= len(rules):
            break
    return anchors if next_rule == len(rules) else None


def scan_candidate_anchors(pages: list[ExtractedPage], rules: list[tuple[str, str, str, re.Pattern[str]]]) -> list[dict[str, Any]]:
    found: list[dict[str, Any]] = []
    seen: set[str] = set()
    for page in pages:
        lines = [clean_line(line) for line in page.text.splitlines() if clean_line(line)]
        for line_no, line in enumerate(lines[:80]):
            if not is_title_anchor_line(line):
                continue
            for subject_code, subject_type, rule_name, pattern in rules:
                if rule_name in seen or not pattern.search(line):
                    continue
                found.append({
                    "pageNo": page.page_no,
                    "lineNo": line_no,
                    "titleAnchor": title_anchor(line),
                    "subjectCode": subject_code,
                    "subjectType": subject_type,
                    "rule": rule_name,
                })
                seen.add(rule_name)
    return found


def is_title_anchor_line(line: str) -> bool:
    if len(line) > 60:
        return False
    if re.search(r"(答案|解析|问题|第\s*\d+\s*题|试题\s*[（(]?\d|[A-H]\s*[.．、])", line):
        return False
    return True


def title_anchor(line: str) -> str:
    return line[:80]


def segmented_plan(document: SourceDocument, scope: DocumentScope, anchors: list[dict[str, Any]]) -> dict[str, Any]:
    segments: list[dict[str, Any]] = []
    last_page = document.pages[-1].page_no if document.pages else 1
    for index, anchor in enumerate(anchors):
        next_anchor = anchors[index + 1] if index + 1 < len(anchors) else None
        segments.append({
            "segmentId": f"S{index + 1:02d}",
            "qualificationCode": scope.qualification_code,
            "subjectCode": anchor["subjectCode"],
            "subjectType": anchor["subjectType"],
            "paperType": scope.paper_type,
            "paperKind": scope.paper_kind,
            "yearNo": scope.year_no,
            "titleAnchor": anchor["titleAnchor"],
            "startPage": anchor["pageNo"],
            "startLine": anchor.get("lineNo", 0),
            "endPage": next_anchor["pageNo"] if next_anchor else last_page,
            "endLineExclusive": next_anchor.get("lineNo") if next_anchor else None,
            "confidence": "HIGH",
            "rule": anchor["rule"],
        })
    return {"status": "SEGMENTED", "segments": segments, "errors": []}


def failed_split_plan(code: str, message: str, detected_anchors: list[dict[str, Any]], required_rules: list[str]) -> dict[str, Any]:
    detected_rules = [anchor["rule"] for anchor in detected_anchors]
    return {
        "status": "FAILED_CLOSED",
        "segments": [],
        "detectedAnchors": detected_anchors,
        "requiredRules": required_rules,
        "missingRules": [rule for rule in required_rules if rule not in detected_rules],
        "errors": [{"code": code, "message": message, "locator": None}],
    }


def segment_to_dict(segment: SegmentScope) -> dict[str, Any]:
    return {
        "segmentId": segment.segment_id,
        "qualificationCode": segment.qualification_code,
        "subjectCode": segment.subject_code,
        "subjectType": segment.subject_type,
        "paperType": segment.paper_type,
        "paperKind": segment.paper_kind,
        "yearNo": segment.year_no,
        "titleAnchor": segment.title_anchor,
        "startPage": segment.start_page,
        "endPage": segment.end_page,
        "confidence": segment.confidence,
        "rule": segment.rule,
    }


def split_plan_errors(path: Path, scope: DocumentScope, split_plan: dict[str, Any]) -> list[dict[str, str]]:
    errors: list[dict[str, str]] = []
    if split_plan["status"] == "FAILED_CLOSED":
        for row in split_plan["errors"]:
            errors.append({"code": row["code"], "message": f"{row['message']} for {path.name}", "locator": row.get("locator")})
    if scope.qualification_code == "UNKNOWN":
        errors.append({"code": "SOURCE_SCOPE_UNKNOWN", "message": "qualification and subject could not be inferred", "locator": None})
    return errors


def build_question_records(document: SourceDocument, scope: DocumentScope, split_plan: dict[str, Any]) -> list[dict[str, Any]]:
    if document.structured_questions:
        records = [
            canonical_structured_question(row, document, index)
            for index, row in enumerate(document.structured_questions, start=1)
        ]
        for record in records:
            finalize_record_quality(record)
        return records
    if split_plan["status"] == "FAILED_CLOSED":
        return []
    records: list[dict[str, Any]] = []
    item_no = 0
    for segment in split_plan["segments"]:
        segment_pages = pages_for_segment(document.pages, segment)
        question_scope = scope_from_segment(segment)
        if question_scope.subject_type == "ESSAY":
            essay_candidate = split_essay_question(segment_pages)
            if essay_candidate:
                item_no += 1
                records.append(canonical_essay_question(
                    essay_candidate, document, question_scope, item_no, segment
                ))
            continue
        if question_scope.subject_type in {"APPLICATION", "CASE"}:
            case_candidates = split_case_questions(segment_pages)
            if case_candidates:
                for candidate in case_candidates:
                    item_no += 1
                    records.append(canonical_case_question(
                        candidate, document, question_scope, item_no, segment
                    ))
                continue
        for candidate in split_questions(segment_pages):
            item_no += 1
            records.append(canonical_question(
                candidate, document, question_scope, item_no, segment
            ))
    for record in records:
        finalize_record_quality(record)
    return records


def pages_for_segment(pages: list[ExtractedPage], segment: dict[str, Any]) -> list[ExtractedPage]:
    selected: list[ExtractedPage] = []
    start_line = segment.get("startLine")
    end_line_exclusive = segment.get("endLineExclusive")
    for page in pages:
        if not segment["startPage"] <= page.page_no <= segment["endPage"]:
            continue
        lines = [clean_line(line) for line in page.text.splitlines() if clean_line(line)]
        lower = start_line if page.page_no == segment["startPage"] and start_line is not None else 0
        upper = end_line_exclusive if page.page_no == segment["endPage"] and end_line_exclusive is not None else len(lines)
        if upper > lower:
            selected.append(ExtractedPage(page.page_no, "\n".join(lines[lower:upper])))
    return selected


def segment_for_page(split_plan: dict[str, Any], page_no: int | None) -> dict[str, Any] | None:
    if page_no is None:
        return split_plan["segments"][0] if split_plan["segments"] else None
    for segment in split_plan["segments"]:
        if segment["startPage"] <= page_no <= segment["endPage"]:
            return segment
    return None


def scope_from_segment(segment: dict[str, Any]) -> DocumentScope:
    return DocumentScope(
        segment["qualificationCode"],
        segment["subjectCode"],
        segment["subjectType"],
        segment["paperType"],
        segment["paperKind"],
        segment["yearNo"],
        segment["confidence"],
    )


def split_questions(pages: list[ExtractedPage]) -> list[dict[str, Any]]:
    candidates: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    for page in pages:
        for raw_line in page.text.splitlines():
            line = clean_line(raw_line)
            if not line:
                continue
            match = QUESTION_START.match(line)
            if match:
                if current:
                    candidates.append(current)
                current = {
                    "questionNo": question_number_from_match(match),
                    "pageNo": page.page_no,
                    "lines": [line],
                }
            elif current:
                current["lines"].append(line)
    if current:
        candidates.append(current)
    if not candidates and any(page.text.strip() for page in pages):
        merged = "\n".join(clean_line(line) for page in pages for line in page.text.splitlines() if clean_line(line))
        if merged:
            candidates.append({
                "questionNo": 1,
                "pageNo": pages[0].page_no if pages else None,
                "lines": [merged],
                "fallbackWholeDocument": True,
            })
    return candidates


def question_number_from_match(match: re.Match[str]) -> int:
    for group_name in ("paren", "ordinal", "plain"):
        value = match.group(group_name)
        if value:
            return int(value)
    raise ValueError("question number regex matched without a number")


def split_case_questions(pages: list[ExtractedPage]) -> list[dict[str, Any]]:
    cases: list[dict[str, Any]] = []
    current_case: dict[str, Any] | None = None
    current_sub: dict[str, Any] | None = None
    in_reference_answers = False
    current_answer_no: int | None = None
    current_answer_lines: list[str] = []

    def ensure_case(page_no: int) -> dict[str, Any]:
        nonlocal current_case
        if current_case is None:
            current_case = {
                "questionNo": len(cases) + 1,
                "pageNo": page_no,
                "materialLines": [],
                "subQuestions": [],
                "answersByNo": {},
            }
        return current_case

    def flush_sub() -> None:
        nonlocal current_sub
        if current_sub is not None and current_case is not None:
            current_case["subQuestions"].append(current_sub)
        current_sub = None

    def flush_answer() -> None:
        nonlocal current_answer_no, current_answer_lines
        if current_case is not None and current_answer_no is not None:
            answer = "\n".join(current_answer_lines).strip()
            answer = re.sub(r"^答\s*[：:]\s*", "", answer)
            current_case["answersByNo"][current_answer_no] = answer
        current_answer_no = None
        current_answer_lines = []

    def flush_case() -> None:
        nonlocal current_case, in_reference_answers
        flush_sub()
        flush_answer()
        if current_case is not None and current_case["subQuestions"]:
            for sub in current_case["subQuestions"]:
                if not sub["answer"]:
                    sub["answer"] = current_case["answersByNo"].get(sub["subQuestionNo"], "")
            current_case.pop("answersByNo", None)
            cases.append(current_case)
        current_case = None
        in_reference_answers = False

    for page in pages:
        for raw_line in page.text.splitlines():
            line = clean_line(raw_line)
            if not line:
                continue
            if CASE_TITLE_LINE.match(line):
                flush_case()
                ensure_case(page.page_no)
                continue
            if REFERENCE_SECTION_LINE.match(line):
                flush_sub()
                flush_answer()
                ensure_case(page.page_no)
                in_reference_answers = True
                continue
            sub_match = CASE_SUBQUESTION_LINE.match(line)
            if sub_match:
                ensure_case(page.page_no)
                if in_reference_answers:
                    flush_answer()
                    current_answer_no = int(sub_match.group(1))
                    answer_text = sub_match.group(3).strip()
                    current_answer_lines = [answer_text] if answer_text else []
                    continue
                flush_sub()
                score = float(sub_match.group(2)) if sub_match.group(2) else None
                if score is not None and score.is_integer():
                    score = int(score)
                current_sub = {
                    "subQuestionNo": int(sub_match.group(1)),
                    "pageNo": page.page_no,
                    "score": score,
                    "lines": [sub_match.group(3).strip()] if sub_match.group(3).strip() else [],
                    "answer": "",
                }
                continue
            if current_sub is not None:
                answer_match = REFERENCE_ANSWER_LINE.match(line)
                if answer_match:
                    current_sub["answer"] = answer_match.group(1).strip()
                else:
                    current_sub["lines"].append(line)
                continue
            if in_reference_answers and current_answer_no is not None:
                current_answer_lines.append(line)
                continue
            case = ensure_case(page.page_no)
            if line not in {"[说明]", "【说明】", "说明"}:
                case["materialLines"].append(line)
    flush_case()
    return cases


def split_essay_question(pages: list[ExtractedPage]) -> dict[str, Any] | None:
    lines = [clean_line(line) for page in pages for line in page.text.splitlines() if clean_line(line)]
    if not lines:
        return None
    title = ""
    requirements: list[str] = []
    reference_parts: list[str] = []
    for line in lines:
        if match := ESSAY_TITLE_LINE.match(line):
            title = match.group(1).strip()
        elif match := ESSAY_REQUIREMENT_LINE.match(line):
            requirements.append(match.group(1).strip())
        elif match := ESSAY_REFERENCE_LINE.match(line):
            reference_parts.append(match.group(1).strip())
    if not title:
        title = next((line for line in lines if line not in {"论文", "论文题"}), lines[0])
    merged = "\n".join(lines)
    word_match = WORD_COUNT_RANGE.search(merged)
    return {
        "questionNo": 1,
        "pageNo": pages[0].page_no if pages else None,
        "title": title,
        "requirements": "\n".join(requirements),
        "minWords": int(word_match.group(1)) if word_match else None,
        "maxWords": int(word_match.group(2)) if word_match else None,
        "reference": "\n".join(reference_parts),
        "rawText": merged,
    }


def clean_line(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\u00a0", " ")).strip()


def canonical_structured_question(
    row: dict[str, Any], document: SourceDocument, item_no: int
) -> dict[str, Any]:
    scope = DocumentScope(
        row["qualificationCode"],
        row["subjectCode"],
        row["subjectType"],
        row["paperType"],
        row["paperKind"],
        row["yearNo"],
        "HIGH" if "UNKNOWN" not in {row["qualificationCode"], row["subjectCode"], row["questionType"]} else "LOW",
    )
    segment = {
        "segmentId": f"SHEET{row['sheetIndex']:02d}",
        "titleAnchor": row["sheetName"],
    }
    question_no = row["rowNo"] - 1
    source_key = row["sourceQuestionId"] or f"ROW-{row['rowNo']}"
    external_key = f"Q_{document.sha256[:12]}_S{row['sheetIndex']:02d}_R{row['rowNo']:06d}_{sha256_text(source_key)[:8]}"
    options = parse_options_text(row["optionsText"])
    question_type = row["questionType"]
    answer_format = "TEXT"
    grading_mode = "MANUAL"
    answer_payload: dict[str, Any] | None = None
    normalized_answer = normalize_answer(row["answerText"]) if row["answerText"] else ""
    if question_type in {"SINGLE_CHOICE", "MULTI_CHOICE"}:
        answer_format, grading_mode = "OPTION_KEYS", "AUTO"
        answer_payload = build_answer_payload([normalized_answer], answer_format) if normalized_answer else None
    elif question_type == "JUDGE":
        answer_format, grading_mode = "BOOLEAN", "AUTO"
        answer_payload = build_answer_payload([normalized_answer], answer_format) if normalized_answer else None
    elif row["answerText"]:
        answer_payload = build_answer_payload([row["answerText"]], answer_format)
    explanation = row["explanation"]
    content_hash = sha256_text(json.dumps({
        "stem": clean_line(row["stem"]),
        "options": options,
        "answer": answer_payload,
        "scope": [scope.qualification_code, scope.subject_code, scope.paper_kind],
    }, ensure_ascii=False, sort_keys=True, separators=(",", ":")))
    record = canonical_record_envelope(
        document,
        scope,
        segment,
        row["sheetIndex"],
        question_no,
        external_key,
        question_type,
        content_hash,
    )
    bindings: list[dict[str, Any]] = []
    if row["chapter"]:
        bindings.append({"bindingType": "CHAPTER", "externalName": row["chapter"], "sortOrder": len(bindings) + 1})
    if row["knowledgePoint"]:
        bindings.append({"bindingType": "KNOWLEDGE_POINT", "externalName": row["knowledgePoint"], "sortOrder": len(bindings) + 1})
    record["sourceTrace"].update({
        "sheetName": row["sheetName"],
        "rowNo": row["rowNo"],
        "sourceQuestionId": row["sourceQuestionId"],
        "sourceUrl": row["sourceUrl"],
        "locator": f"sha256:{document.sha256}:sheet:{row['sheetName']}:row:{row['rowNo']}",
    })
    record["question"].update({
        "title": source_key,
        "stems": [{
            "stemKey": "S1",
            "stemType": "MAIN_STEM",
            "content": row["stem"],
            "resourceId": None,
            "sortOrder": 1,
        }],
        "subQuestions": [{
            "subQuestionNo": 1,
            "subQuestionType": question_type,
            "content": row["stem"],
            "score": None,
            "answerFormat": answer_format,
            "gradingMode": grading_mode,
            "sortOrder": 1,
            "options": options,
            "answers": [answer_payload] if answer_payload else [],
            "solutions": [{
                "solutionType": "STANDARD",
                "content": explanation,
                "sortOrder": 1,
            }] if explanation else [],
            "scoringRules": [{
                "ruleType": "AUTO" if grading_mode == "AUTO" else "MANUAL",
                "ruleJson": {"mode": "EXACT" if grading_mode == "AUTO" else "MANUAL", "tolerance": None, "caseSensitive": None},
                "sortOrder": 1,
            }],
            "scoringPoints": [],
            "knowledgeBindings": bindings,
        }],
        "knowledgeBindings": bindings,
        "source": {
            "sourceType": "SOURCE_DOCUMENT_IMPORT",
            "sourceRef": row["sourceUrl"] or f"sha256:{document.sha256}",
            "sourceName": row["sourceName"],
            "syllabusBasis": "UNVERIFIED_PREFLIGHT",
            "copyrightStatus": "UNVERIFIED",
            "usageRestriction": "local development/internal acceptance only; publication prohibited",
            "capturedAt": row["capturedAt"],
        },
    })
    record["quality"] = {
        "parserConfidence": scope.confidence,
        "needsManualReview": scope.confidence == "LOW" or not row["answerText"],
    }
    return record


def parse_options_text(value: str) -> list[dict[str, Any]]:
    if not value:
        return []
    pattern = re.compile(r"(?m)^\s*([A-H])\s*[.．、)]\s*")
    matches = list(pattern.finditer(value))
    options: list[dict[str, Any]] = []
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(value)
        content = value[match.end():end].strip()
        options.append({
            "optionKey": match.group(1).upper(),
            "content": content,
            "sortOrder": index + 1,
        })
    return options


def parse_option_line(value: str) -> list[dict[str, Any]]:
    pattern = re.compile(r"(?<!\S)([A-H])\s*[.．、)]\s*", re.IGNORECASE)
    matches = list(pattern.finditer(value))
    if not matches or value[:matches[0].start()].strip():
        return []
    options: list[dict[str, Any]] = []
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(value)
        options.append({
            "optionKey": match.group(1).upper(),
            "content": value[match.end():end].strip(),
            "sortOrder": index + 1,
        })
    return options


def canonical_case_question(
    candidate: dict[str, Any],
    document: SourceDocument,
    scope: DocumentScope,
    item_no: int,
    segment: dict[str, Any],
) -> dict[str, Any]:
    question_no = candidate["questionNo"]
    page_no = candidate.get("pageNo")
    material = "\n".join(candidate.get("materialLines") or []).strip()
    sub_questions: list[dict[str, Any]] = []
    for sort_order, sub in enumerate(candidate["subQuestions"], start=1):
        content = "\n".join(sub.get("lines") or []).strip()
        answer = str(sub.get("answer") or "").strip()
        answer_payload = {
            "answerKind": "TEXT",
            "optionKeys": None,
            "booleanValue": None,
            "textValue": answer,
            "numberValue": None,
            "sortOrder": 1,
        } if answer else None
        sub_questions.append({
            "subQuestionNo": sub["subQuestionNo"],
            "subQuestionType": "SHORT_ANSWER",
            "content": content,
            "score": sub.get("score"),
            "answerFormat": "TEXT",
            "gradingMode": "MANUAL",
            "sortOrder": sort_order,
            "options": [],
            "answers": [answer_payload] if answer_payload else [],
            "solutions": [{
                "solutionType": "REFERENCE_ANSWER",
                "content": answer,
                "sortOrder": 1,
            }] if answer else [],
            "scoringRules": [{
                "ruleType": "MANUAL",
                "ruleJson": {"mode": "MANUAL", "tolerance": None, "caseSensitive": None},
                "sortOrder": 1,
            }],
            "scoringPoints": [{
                "pointCode": "P1",
                "pointText": answer,
                "score": sub.get("score"),
                "sortOrder": 1,
            }] if answer and sub.get("score") is not None else [],
            "knowledgeBindings": [],
        })
    external_key = external_question_key(document.sha256, page_no, question_no, item_no)
    content_hash = sha256_text(json.dumps({
        "material": material,
        "subQuestions": sub_questions,
        "scope": [scope.qualification_code, scope.subject_code, scope.paper_kind],
    }, ensure_ascii=False, sort_keys=True, separators=(",", ":")))
    record = canonical_record_envelope(
        document, scope, segment, page_no, question_no, external_key, "CASE", content_hash
    )
    record["question"].update({
        "title": f"{scope.paper_kind} case {question_no}",
        "stems": [{
            "stemKey": "M1",
            "stemType": "CASE_MATERIAL",
            "content": material,
            "resourceId": None,
            "sortOrder": 1,
        }],
        "subQuestions": sub_questions,
    })
    record["quality"] = {
        "parserConfidence": "HIGH" if material and all(sub["answers"] for sub in sub_questions) else "MEDIUM",
        "needsManualReview": not material or any(not sub["answers"] for sub in sub_questions),
    }
    return record


def canonical_essay_question(
    candidate: dict[str, Any],
    document: SourceDocument,
    scope: DocumentScope,
    item_no: int,
    segment: dict[str, Any],
) -> dict[str, Any]:
    question_no = candidate["questionNo"]
    page_no = candidate.get("pageNo")
    title = candidate["title"]
    requirements = candidate.get("requirements") or candidate.get("rawText") or ""
    reference = candidate.get("reference") or ""
    external_key = external_question_key(document.sha256, page_no, question_no, item_no)
    content_hash = sha256_text(json.dumps({
        "title": title,
        "requirements": requirements,
        "minWords": candidate.get("minWords"),
        "maxWords": candidate.get("maxWords"),
        "reference": reference,
        "scope": [scope.qualification_code, scope.subject_code, scope.paper_kind],
    }, ensure_ascii=False, sort_keys=True, separators=(",", ":")))
    record = canonical_record_envelope(
        document, scope, segment, page_no, question_no, external_key, "ESSAY", content_hash
    )
    answer_payload = {
        "answerKind": "TEXT",
        "optionKeys": None,
        "booleanValue": None,
        "textValue": reference,
        "numberValue": None,
        "sortOrder": 1,
    } if reference else None
    record["question"].update({
        "title": title,
        "essayRequirements": {
            "writingRequirements": requirements,
            "minWords": candidate.get("minWords"),
            "maxWords": candidate.get("maxWords"),
        },
        "stems": [{
            "stemKey": "S1",
            "stemType": "MAIN_STEM",
            "content": title,
            "resourceId": None,
            "sortOrder": 1,
        }],
        "subQuestions": [{
            "subQuestionNo": 1,
            "subQuestionType": "ESSAY",
            "content": requirements,
            "score": None,
            "answerFormat": "LONG_TEXT",
            "gradingMode": "MANUAL",
            "sortOrder": 1,
            "options": [],
            "answers": [answer_payload] if answer_payload else [],
            "solutions": [{
                "solutionType": "REFERENCE_STRUCTURE",
                "content": reference,
                "sortOrder": 1,
            }] if reference else [],
            "scoringRules": [{
                "ruleType": "MANUAL",
                "ruleJson": {"mode": "MANUAL", "tolerance": None, "caseSensitive": None},
                "sortOrder": 1,
            }],
            "scoringPoints": [],
            "knowledgeBindings": [],
        }],
    })
    record["quality"] = {
        "parserConfidence": "HIGH" if title and requirements else "MEDIUM",
        "needsManualReview": not title or not requirements or not reference,
    }
    return record


def canonical_record_envelope(
    document: SourceDocument,
    scope: DocumentScope,
    segment: dict[str, Any],
    page_no: int | None,
    question_no: int,
    external_key: str,
    question_type: str,
    content_hash: str,
) -> dict[str, Any]:
    return {
        "schemaVersion": SCHEMA_VERSION,
        "recordType": "QUESTION_PREFLIGHT",
        "rightsStatus": RIGHTS_STATUS,
        "importStatus": IMPORT_STATUS,
        "qualificationCode": scope.qualification_code,
        "subjectCode": scope.subject_code,
        "subjectType": scope.subject_type,
        "paper": {
            "paperKey": paper_key(document.sha256, scope, segment["segmentId"]),
            "paperType": scope.paper_type,
            "paperKind": scope.paper_kind,
            "yearNo": scope.year_no,
            "sourceDocumentSha256": document.sha256,
            "sourceFileName": document.path.name,
        },
        "sourceTrace": {
            "documentSha256": document.sha256,
            "sourcePath": str(document.path),
            "segmentId": segment["segmentId"],
            "segmentTitleAnchor": segment["titleAnchor"],
            "pageNo": page_no,
            "questionNo": question_no,
            "locator": f"sha256:{document.sha256}:segment:{segment['segmentId']}:page:{page_no}:question:{question_no}",
        },
        "question": {
            "externalKey": external_key,
            "questionType": question_type,
            "title": "",
            "yearNo": scope.year_no,
            "difficulty": None,
            "source": {
                "sourceType": "SOURCE_DOCUMENT_IMPORT",
                "sourceRef": f"sha256:{document.sha256}",
                "syllabusBasis": "UNVERIFIED_PREFLIGHT",
                "copyrightStatus": "UNVERIFIED",
                "usageRestriction": "local development/internal acceptance only; publication prohibited",
            },
            "stems": [],
            "subQuestions": [],
            "knowledgeBindings": [],
            "contentHash": content_hash,
        },
        "quality": {},
    }


def canonical_question(candidate: dict[str, Any], document: SourceDocument, scope: DocumentScope, item_no: int, segment: dict[str, Any]) -> dict[str, Any]:
    lines = candidate["lines"]
    question_no = candidate["questionNo"]
    page_no = candidate.get("pageNo")
    options: list[dict[str, Any]] = []
    stem_lines: list[str] = []
    answers: list[str] = []
    solution_lines: list[str] = []
    in_solution = False
    for line in lines:
        line_options = parse_option_line(line)
        answer_match = ANSWER_LINE.search(line)
        explanation_match = EXPLANATION_LINE.match(line)
        if explanation_match:
            if explanation_match.group(1):
                solution_lines.append(explanation_match.group(1).strip())
            in_solution = True
        elif in_solution:
            solution_lines.append(line)
        elif line_options:
            for option in line_options:
                option["sortOrder"] = len(options) + 1
                options.append(option)
            in_solution = False
        elif answer_match:
            answers.append(normalize_answer(answer_match.group(1)))
            in_solution = False
        else:
            stem_lines.append(line)
    stem_text = "\n".join(stem_lines).strip() or "\n".join(lines).strip()
    question_type, sub_type, answer_format, grading_mode = infer_question_type(scope, options, answers, stem_text)
    external_key = external_question_key(document.sha256, page_no, question_no, item_no)
    answer_payload = build_answer_payload(answers, answer_format)
    solution = "\n".join(solution_lines).strip()
    content_hash = sha256_text(json.dumps({
        "stem": stem_text,
        "options": options,
        "answer": answer_payload,
        "scope": [scope.qualification_code, scope.subject_code, scope.paper_kind],
    }, ensure_ascii=False, sort_keys=True, separators=(",", ":")))
    return {
        "schemaVersion": SCHEMA_VERSION,
        "recordType": "QUESTION_PREFLIGHT",
        "rightsStatus": RIGHTS_STATUS,
        "importStatus": IMPORT_STATUS,
        "qualificationCode": scope.qualification_code,
        "subjectCode": scope.subject_code,
        "subjectType": scope.subject_type,
        "paper": {
            "paperKey": paper_key(document.sha256, scope, segment["segmentId"]),
            "paperType": scope.paper_type,
            "paperKind": scope.paper_kind,
            "yearNo": scope.year_no,
            "sourceDocumentSha256": document.sha256,
            "sourceFileName": document.path.name,
        },
        "sourceTrace": {
            "documentSha256": document.sha256,
            "sourcePath": str(document.path),
            "segmentId": segment["segmentId"],
            "segmentTitleAnchor": segment["titleAnchor"],
            "pageNo": page_no,
            "questionNo": question_no,
            "locator": f"sha256:{document.sha256}:segment:{segment['segmentId']}:page:{page_no}:question:{question_no}",
        },
        "question": {
            "externalKey": external_key,
            "questionType": question_type,
            "title": f"{scope.paper_kind} {question_no}",
            "yearNo": scope.year_no,
            "difficulty": None,
            "source": {
                "sourceType": "SOURCE_DOCUMENT_IMPORT",
                "sourceRef": f"sha256:{document.sha256}",
                "syllabusBasis": "UNVERIFIED_PREFLIGHT",
                "copyrightStatus": "UNVERIFIED",
                "usageRestriction": "local development/internal acceptance only; publication prohibited",
            },
            "stems": [{
                "stemKey": "S1" if question_type != "CASE" else "M1",
                "stemType": "MAIN_STEM" if question_type != "CASE" else "CASE_MATERIAL",
                "content": stem_text,
                "resourceId": None,
                "sortOrder": 1,
            }],
            "subQuestions": [{
                "subQuestionNo": 1,
                "subQuestionType": sub_type,
                "content": stem_text if question_type != "CASE" else "See case material.",
                "score": None,
                "answerFormat": answer_format,
                "gradingMode": grading_mode,
                "sortOrder": 1,
                "options": options,
                "answers": [answer_payload] if answer_payload else [],
                "solutions": [{"solutionType": "REFERENCE_ANSWER", "content": solution, "sortOrder": 1}] if solution else [],
                "scoringRules": [{
                    "ruleType": "AUTO" if grading_mode == "AUTO" else "MANUAL",
                    "ruleJson": {"mode": "EXACT" if grading_mode == "AUTO" else "MANUAL", "tolerance": None, "caseSensitive": None},
                    "sortOrder": 1,
                }],
                "scoringPoints": [],
                "knowledgeBindings": [],
            }],
            "knowledgeBindings": [],
            "contentHash": content_hash,
        },
        "quality": {
            "parserConfidence": parser_confidence(options, answers, stem_text, scope),
            "needsManualReview": scope.confidence == "LOW" or not stem_text or not answers,
            "fallbackWholeDocument": bool(candidate.get("fallbackWholeDocument")),
        },
    }


def infer_question_type(scope: DocumentScope, options: list[dict[str, Any]], answers: list[str], stem_text: str) -> tuple[str, str, str, str]:
    if scope.subject_type in {"APPLICATION", "CASE"} or (not options and len(stem_text) > 300):
        return "CASE", "SHORT_ANSWER", "TEXT", "MANUAL"
    if answers and answers[0] in {"TRUE", "FALSE"}:
        return "JUDGE", "JUDGE", "BOOLEAN", "AUTO"
    if len(options) >= 2:
        if answers and len(answers[0]) > 1:
            return "MULTI_CHOICE", "MULTI_CHOICE", "OPTION_KEYS", "AUTO"
        return "SINGLE_CHOICE", "SINGLE_CHOICE", "OPTION_KEYS", "AUTO"
    return "SHORT_ANSWER", "SHORT_ANSWER", "TEXT", "MANUAL"


def normalize_answer(value: str) -> str:
    text = value.strip().upper()
    if text in {"正确", "对", "TRUE"}:
        return "TRUE"
    if text in {"错误", "错", "FALSE"}:
        return "FALSE"
    return re.sub(r"[^A-H0-9.]", "", text)


def build_answer_payload(answers: list[str], answer_format: str) -> dict[str, Any] | None:
    if not answers:
        return None
    answer = answers[0]
    if answer_format == "OPTION_KEYS":
        keys = sorted(set(answer))
        return {"answerKind": "OPTION_KEYS", "optionKeys": keys, "booleanValue": None, "textValue": None, "numberValue": None, "sortOrder": 1}
    if answer_format == "BOOLEAN":
        return {"answerKind": "BOOLEAN", "optionKeys": None, "booleanValue": answer == "TRUE", "textValue": None, "numberValue": None, "sortOrder": 1}
    if answer_format == "NUMBER":
        return {"answerKind": "NUMBER", "optionKeys": None, "booleanValue": None, "textValue": None, "numberValue": float(answer), "sortOrder": 1}
    return {"answerKind": "TEXT", "optionKeys": None, "booleanValue": None, "textValue": answer, "numberValue": None, "sortOrder": 1}


def parser_confidence(options: list[dict[str, Any]], answers: list[str], stem_text: str, scope: DocumentScope) -> str:
    if scope.confidence == "LOW":
        return "LOW"
    if options and answers and stem_text:
        return "HIGH"
    if stem_text:
        return "MEDIUM"
    return "LOW"


def external_question_key(document_sha: str, page_no: int | None, question_no: int, item_no: int) -> str:
    page = page_no if page_no is not None else 0
    return f"Q_{document_sha[:12]}_P{page:03d}_N{question_no:03d}_I{item_no:04d}"


def paper_key(document_sha: str, scope: DocumentScope, segment_id: str) -> str:
    year = scope.year_no or "UNKNOWN"
    return f"{scope.qualification_code}:{scope.subject_code}:{scope.paper_kind}:{year}:{document_sha[:12]}:{segment_id}"


def duplicate_rows(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    first_by_hash: dict[str, dict[str, Any]] = {}
    duplicates: list[dict[str, Any]] = []
    for record in records:
        content_hash = record["question"]["contentHash"]
        first = first_by_hash.get(content_hash)
        if first is None:
            first_by_hash[content_hash] = record
            continue
        duplicates.append({
            "contentHash": content_hash,
            "firstExternalKey": first["question"]["externalKey"],
            "duplicateExternalKey": record["question"]["externalKey"],
            "firstSource": first["sourceTrace"]["locator"],
            "duplicateSource": record["sourceTrace"]["locator"],
        })
    return duplicates


def validate_record(record: dict[str, Any]) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    if record["qualificationCode"] in {"", "UNKNOWN"}:
        issues.append(validation_issue(
            "SOURCE_SCOPE_UNKNOWN", "qualificationCode", "qualification could not be mapped"
        ))
    if record["subjectCode"] in {"", "UNKNOWN"} or record["subjectType"] in {"", "UNKNOWN", "MIXED"}:
        issues.append(validation_issue(
            "SUBJECT_SCOPE_UNKNOWN", "subjectCode", "subject could not be mapped to one exam paper"
        ))
    if record["question"]["questionType"] in {"", "UNKNOWN"}:
        issues.append(validation_issue(
            "QUESTION_TYPE_UNKNOWN", "question.questionType", "question type could not be mapped"
        ))
    if record["paper"]["paperKind"] in {"", "UNKNOWN"}:
        issues.append(validation_issue(
            "PAPER_KIND_UNKNOWN", "paper.paperKind", "source is neither verified REAL_PAPER nor MOCK_PAPER"
        ))
    if record["paper"].get("yearNo") is None:
        issues.append(validation_issue("YEAR_UNKNOWN", "paper.yearNo", "exam year is required for import"))
    if record["quality"].get("fallbackWholeDocument"):
        issues.append(validation_issue(
            "DOCUMENT_FALLBACK_UNRESOLVED",
            "question.stems",
            "document had no reliable question boundaries and was retained only as an unresolved fallback",
        ))
    if (
        record["subjectType"] == "COMPREHENSIVE"
        and record["question"]["questionType"] not in {"SINGLE_CHOICE", "MULTI_CHOICE", "JUDGE"}
    ):
        issues.append(validation_issue(
            "COMPREHENSIVE_TYPE_INVALID",
            "question.questionType",
            "comprehensive/basic knowledge records must be choice or judgement questions",
        ))
    stems = record["question"].get("stems") or []
    if not stems or not any(str(stem.get("content") or "").strip() for stem in stems):
        issues.append(validation_issue("STEM_MISSING", "question.stems", "question stem is required"))

    sub_questions = record["question"].get("subQuestions") or []
    if record["question"]["questionType"] == "CASE" and not 2 <= len(sub_questions) <= 4:
        issues.append(validation_issue(
            "CASE_SUBQUESTION_COUNT_INVALID",
            "question.subQuestions",
            "case questions must contain two to four ordered sub-questions",
        ))

    for sub_index, sub_question in enumerate(sub_questions, start=1):
        prefix = f"question.subQuestions[{sub_index}]"
        question_type = record["question"]["questionType"]
        options = sub_question.get("options") or []
        answers = sub_question.get("answers") or []
        solutions = sub_question.get("solutions") or []
        if question_type in {"SINGLE_CHOICE", "MULTI_CHOICE"}:
            option_keys = [str(option.get("optionKey") or "").upper() for option in options]
            expected = [chr(ord("A") + index) for index in range(len(option_keys))]
            if len(option_keys) < 2 or option_keys != expected:
                issues.append(validation_issue(
                    "INVALID_OPTION_SEQUENCE",
                    f"{prefix}.options",
                    f"option keys must be contiguous from A; got {option_keys}",
                ))
            if not answers:
                issues.append(validation_issue("ANSWER_MISSING", f"{prefix}.answers", "answer is required"))
            else:
                answer_keys = answers[0].get("optionKeys") or []
                if not answer_keys or any(key not in option_keys for key in answer_keys):
                    issues.append(validation_issue(
                        "ANSWER_OPTION_REFERENCE_INVALID",
                        f"{prefix}.answers[1].optionKeys",
                        f"answer keys {answer_keys} must reference existing options",
                    ))
                if question_type == "SINGLE_CHOICE" and len(answer_keys) != 1:
                    issues.append(validation_issue(
                        "SINGLE_CHOICE_ANSWER_CARDINALITY",
                        f"{prefix}.answers[1].optionKeys",
                        "single-choice answer must contain exactly one option key",
                    ))
        elif question_type == "CASE":
            score = decimal_or_none(sub_question.get("score"))
            scoring_points = sub_question.get("scoringPoints") or []
            if not answers:
                issues.append(validation_issue(
                    "ANSWER_MISSING", f"{prefix}.answers", "case reference answer is required"
                ))
            if score is None or score <= 0:
                issues.append(validation_issue(
                    "CASE_SCORE_MISSING", f"{prefix}.score", "case sub-question score must be positive"
                ))
            if not scoring_points:
                issues.append(validation_issue(
                    "CASE_SCORING_POINTS_MISSING",
                    f"{prefix}.scoringPoints",
                    "case sub-question requires auditable scoring points",
                ))
            else:
                point_scores = [decimal_or_none(point.get("score")) for point in scoring_points]
                if (
                    score is None
                    or any(point_score is None or point_score <= 0 for point_score in point_scores)
                    or sum((point_score for point_score in point_scores if point_score is not None), Decimal("0")) != score
                ):
                    issues.append(validation_issue(
                        "CASE_SCORE_CLOSURE_INVALID",
                        f"{prefix}.scoringPoints",
                        "scoring point scores must be positive and close exactly to the sub-question score",
                    ))
        elif not answers and not solutions:
            issues.append(validation_issue("ANSWER_MISSING", f"{prefix}.answers", "answer or reference solution is required"))
    return issues


def decimal_or_none(value: Any) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None


def finalize_record_quality(record: dict[str, Any]) -> list[dict[str, str]]:
    issues = validate_record(record)
    record["quality"]["validationErrors"] = issues
    record["quality"]["importEligible"] = not issues
    record["quality"]["needsManualReview"] = bool(
        record["quality"].get("needsManualReview") or issues
    )
    return issues


def validation_issue(code: str, field: str, message: str) -> dict[str, str]:
    return {"code": code, "field": field, "message": message}


def unresolved_rows_for_record(record: dict[str, Any], issues: list[dict[str, str]]) -> list[dict[str, Any]]:
    trace = record["sourceTrace"]
    return [
        unresolved_row(
            Path(trace["sourcePath"]),
            trace["documentSha256"],
            trace["locator"],
            record["question"]["externalKey"],
            issue["code"],
            issue["field"],
            issue["message"],
        )
        for issue in issues
    ]


def unresolved_row(
    path: Path,
    document_sha: str,
    locator: str | None,
    external_key: str,
    code: str,
    field: str,
    message: str,
) -> dict[str, Any]:
    return {
        "sourcePath": str(path),
        "documentSha256": document_sha,
        "locator": locator or "",
        "externalKey": external_key,
        "code": code,
        "field": field,
        "message": message,
        "blocking": "true",
    }


def stable_unique_unresolved(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    seen: set[tuple[str, str, str, str]] = set()
    unique: list[dict[str, Any]] = []
    for row in rows:
        key = (row["sourcePath"], row["locator"], row["externalKey"], row["code"])
        if key not in seen:
            seen.add(key)
            unique.append(row)
    return unique


def build_manifest(
    files: list[Path],
    documents: list[dict[str, Any]],
    records: list[dict[str, Any]],
    errors: list[dict[str, Any]],
    duplicates: list[dict[str, Any]],
    assets: list[dict[str, Any]],
    unresolved: list[dict[str, Any]],
) -> dict[str, Any]:
    by_kind: dict[str, int] = {}
    by_scope: dict[str, int] = {}
    for record in records:
        by_kind[record["paper"]["paperKind"]] = by_kind.get(record["paper"]["paperKind"], 0) + 1
        scope = f"{record['qualificationCode']}->{record['subjectCode']}"
        by_scope[scope] = by_scope.get(scope, 0) + 1
    return {
        "schemaVersion": SCHEMA_VERSION,
        "tool": "question-ingestion-preflight",
        "toolVersion": "0.3.0",
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "dryRun": True,
        "databaseWrites": False,
        "publish": False,
        "documentCount": len(files),
        "questionCount": len(records),
        "errorCount": len(errors),
        "duplicateCount": len(duplicates),
        "assetCount": len(assets),
        "unresolvedCount": len(unresolved),
        "eligibleQuestionCount": sum(1 for record in records if record["quality"].get("importEligible")),
        "ineligibleQuestionCount": sum(1 for record in records if not record["quality"].get("importEligible")),
        "rightsStatus": RIGHTS_STATUS,
        "importStatus": IMPORT_STATUS,
        "paperKindCounts": by_kind,
        "scopeCounts": by_scope,
        "documents": documents,
    }


def document_manifest(path: Path, document: SourceDocument, scope: DocumentScope, split_plan: dict[str, Any], question_count: int, error_count: int) -> dict[str, Any]:
    return {
        "path": str(path),
        "fileName": path.name,
        "sha256": document.sha256,
        "sizeBytes": document.size_bytes,
        "pageCount": len(document.pages),
        "assetCount": len(document.assets),
        "qualificationCode": scope.qualification_code,
        "subjectCode": scope.subject_code,
        "subjectType": scope.subject_type,
        "paperKind": scope.paper_kind,
        "paperType": scope.paper_type,
        "yearNo": scope.year_no,
        "scopeConfidence": scope.confidence,
        "splitPlan": split_plan,
        "questionCount": question_count,
        "errorCount": error_count,
        "status": "PREFLIGHTED",
    }


def error_row(path: Path, sha: str, code: str, message: str, locator: str | None) -> dict[str, Any]:
    return {"sourcePath": str(path), "documentSha256": sha, "code": code, "message": message, "locator": locator or ""}


def asset_row(path: Path, sha: str, asset_key: str, asset_type: str, page_no: int | None, relationship: str, bytes_value: str) -> dict[str, Any]:
    return {
        "sourcePath": str(path),
        "documentSha256": sha,
        "assetKey": asset_key,
        "assetType": asset_type,
        "pageNo": "" if page_no is None else page_no,
        "relationship": relationship,
        "bytes": bytes_value,
    }


def write_jsonl(path: Path, records: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n")


def write_json(path: Path, data: dict[str, Any]) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def write_quality_report(path: Path, manifest: dict[str, Any], documents: list[dict[str, Any]], errors: list[dict[str, Any]]) -> None:
    lines = [
        "# Question Ingestion Preflight Quality Report",
        "",
        f"- Generated at: `{manifest['generatedAt']}`",
        f"- Documents: `{manifest['documentCount']}`",
        f"- Questions: `{manifest['questionCount']}`",
        f"- Eligible questions: `{manifest['eligibleQuestionCount']}`",
        f"- Ineligible questions: `{manifest['ineligibleQuestionCount']}`",
        f"- Unresolved items: `{manifest['unresolvedCount']}`",
        f"- Errors: `{manifest['errorCount']}`",
        f"- Duplicates: `{manifest['duplicateCount']}`",
        f"- Assets: `{manifest['assetCount']}`",
        f"- Database writes: `{manifest['databaseWrites']}`",
        f"- Publish: `{manifest['publish']}`",
        "",
        "## Scope Counts",
        "",
    ]
    for key, count in sorted(manifest["scopeCounts"].items()):
        lines.append(f"- `{key}`: `{count}`")
    lines.extend(["", "## Documents", ""])
    for doc in documents:
        split_plan = doc.get("splitPlan", {})
        split_status = split_plan.get("status", "UNKNOWN")
        lines.append(f"- `{doc.get('fileName', doc.get('path'))}`: questions=`{doc.get('questionCount', 0)}`, errors=`{doc.get('errorCount', 0)}`, scope=`{doc.get('qualificationCode', 'UNKNOWN')}->{doc.get('subjectCode', 'UNKNOWN')}`, splitPlan=`{split_status}`")
        for segment in split_plan.get("segments", []):
            lines.append(f"  - segment `{segment['segmentId']}`: `{segment['subjectCode']}` pages=`{segment['startPage']}-{segment['endPage']}` anchor=`{segment['titleAnchor']}`")
    if errors:
        lines.extend(["", "## Error Codes", ""])
        counts: dict[str, int] = {}
        for row in errors:
            counts[row["code"]] = counts.get(row["code"], 0) + 1
        for code, count in sorted(counts.items()):
            lines.append(f"- `{code}`: `{count}`")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def sha256_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


if __name__ == "__main__":
    sys.exit(main())

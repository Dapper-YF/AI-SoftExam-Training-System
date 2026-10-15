from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

import openpyxl
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import soft_exam_xlsx_preflight as soft_exam  # noqa: E402


def test_missing_composite_options_are_explicit_and_fully_tracked(
    tmp_path: Path,
) -> None:
    source = write_workbook(
        tmp_path / "source.xlsx",
        [
            source_row(
                "42",
                "共享题干中有(1)和(2)两个空(1)",
                option_group("甲", "乙", "丙", "丁"),
                "A、B",
            )
        ],
    )
    output = tmp_path / "output"

    manifest = run_preflight(source, output, expected=1)
    record = read_jsonl(output / "canonical-questions.jsonl")[0]
    mapping = read_csv(output / "mapping.csv")[0]
    anomalies = read_csv(output / "composite-anomalies.csv")

    assert manifest["sourceTrackingRecordCount"] == 1
    assert manifest["usableCanonicalQuestionCount"] == 0
    assert manifest["exactDuplicateMappingCount"] == 0
    assert manifest["explicitAnomalySourceCount"] == 1
    assert manifest["countReconciliation"]["components"]["sum"] == 1
    assert manifest["errorCodeCounts"]["MISSING_COMPOSITE_OPTIONS"] == 1
    assert record["question"]["frozenAnswer"]["normalized"] == "A、B"
    assert record["question"]["compositeOptionState"]["missingGroupNos"] == [2]
    assert record["quality"]["importEligible"] is False
    assert mapping["status"] == "SOURCE_ERROR"
    assert mapping["sourceRow"] == "2"
    assert mapping["missingCompositeGroupNos"] == "2"
    assert anomalies[0]["status"] == "MISSING_COMPOSITE_OPTIONS"
    assert json.loads(anomalies[0]["presentOptionGroupsJson"])[0][0][
        "optionKey"
    ] == "A"


def test_same_workbook_exact_match_recovers_composite_then_deduplicates(
    tmp_path: Path,
) -> None:
    first_group = option_group("甲", "乙", "丙", "丁")
    second_group = option_group("一", "二", "三", "四")
    source = write_workbook(
        tmp_path / "source.xlsx",
        [
            source_row(
                "42",
                "共享题干中有(1)和(2)两个空(1)",
                first_group,
                "A、B",
            ),
            source_row(
                "43",
                "共享题干中有(1)和(2)两个空(1)",
                f"{first_group}\n{second_group}",
                "A、B",
            ),
        ],
    )
    output = tmp_path / "output"

    manifest = run_preflight(source, output, expected=2)
    records = read_jsonl(output / "canonical-questions.jsonl")
    mappings = read_csv(output / "mapping.csv")
    recovery = read_csv(output / "recovery-evidence.csv")

    assert manifest["recovery"]["sameWorkbookCompleteCandidateCount"] == 1
    assert manifest["recovery"]["recoveredCompositeCount"] == 1
    assert manifest["usableCanonicalQuestionCount"] == 1
    assert manifest["exactDuplicateMappingCount"] == 1
    assert manifest["explicitAnomalySourceCount"] == 0
    assert manifest["countReconciliation"]["components"]["sum"] == 2
    assert records[0]["sourceTrace"]["recoveredFromSourceRow"] == 3
    assert len(records[0]["question"]["subQuestions"]) == 2
    assert records[0]["quality"]["importEligible"] is True
    assert mappings[1]["status"] == "DUPLICATE_SOURCE"
    assert recovery[0]["recoveryMethod"] == "SAME_XLSX_EXACT"
    assert recovery[0]["recoveredFromSourceRow"] == "3"


def test_extra_composite_groups_are_not_reported_as_missing(
    tmp_path: Path,
) -> None:
    groups = "\n".join(
        [
            option_group("甲", "乙", "丙", "丁"),
            option_group("一", "二", "三", "四"),
            option_group("壹", "贰", "叁", "肆"),
        ]
    )
    source = write_workbook(
        tmp_path / "source.xlsx",
        [source_row("42", "两个空的题干(1)", groups, "A、B")],
    )
    output = tmp_path / "output"

    manifest = run_preflight(source, output, expected=1)

    assert manifest["errorCodeCounts"] == {
        "COMPOSITE_OPTION_GROUP_COUNT_MISMATCH": 1
    }
    assert "MISSING_COMPOSITE_OPTIONS" not in manifest["errorCodeCounts"]
    assert read_csv(output / "composite-anomalies.csv")[0]["status"] == (
        "COMPOSITE_OPTION_GROUP_COUNT_MISMATCH"
    )


def test_local_docx_exact_text_recovery_records_file_anchor(
    tmp_path: Path,
) -> None:
    stem_text = (
        "本地文档中的共享题干用于验证确定性文本匹配，"
        "题目包含(1)和(2)两个空(1)"
    )
    first_group = option_group("甲", "乙", "丙", "丁")
    second_group = option_group("一", "二", "三", "四")
    source = write_workbook(
        tmp_path / "source.xlsx",
        [source_row("42", stem_text, first_group, "A、B")],
    )
    local_document = tmp_path / "local-corpus.docx"
    document = Document()
    document.add_paragraph(stem_text)
    for line in f"{first_group}\n{second_group}".splitlines():
        document.add_paragraph(line)
    document.save(local_document)
    output = tmp_path / "output"

    manifest = soft_exam.run_streaming_preflight(
        source,
        output,
        soft_exam.default_taxonomy_manifest(),
        expected_question_count=1,
        chunk_size=1,
        local_corpus_paths=[local_document],
    )
    record = read_jsonl(output / "canonical-questions.jsonl")[0]
    recovery = read_csv(output / "recovery-evidence.csv")[0]

    assert manifest["recovery"]["recoveredCompositeCount"] == 1
    assert manifest["recovery"]["methods"] == {"LOCAL_DOCX_EXACT_TEXT": 1}
    assert manifest["usableCanonicalQuestionCount"] == 1
    assert record["quality"]["importEligible"] is True
    assert record["sourceTrace"]["localRecovery"]["sourcePath"] == str(
        local_document.resolve()
    )
    assert recovery["localSourceAnchor"].startswith("paragraph:1-")
    assert len(recovery["localSourceSha256"]) == 64


def run_preflight(source: Path, output: Path, *, expected: int) -> dict[str, object]:
    return soft_exam.run_streaming_preflight(
        source,
        output,
        soft_exam.default_taxonomy_manifest(),
        expected_question_count=expected,
        chunk_size=1,
    )


def write_workbook(path: Path, rows: list[list[str]]) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "软考题目"
    sheet.append(list(soft_exam.EXPECTED_HEADERS))
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return path


def source_row(
    source_question_id: str, stem: str, options: str, answer: str
) -> list[str]:
    return [
        "本地测试题库",
        "软件设计师",
        "综合知识历年真题",
        "",
        source_question_id,
        "single_choice",
        stem,
        options,
        answer,
        "",
        "",
        "",
        "",
        "2026-07-16T00:00:00",
    ]


def option_group(*values: str) -> str:
    return "\n".join(
        f"{chr(ord('A') + index)}. {value}"
        for index, value in enumerate(values)
    )


def read_jsonl(path: Path) -> list[dict[str, object]]:
    return [
        json.loads(line)
        for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))
